#!/usr/bin/env python3
"""Build a traceable, read-only CCMA workbook from a preserved CSV source."""

from __future__ import annotations

import csv
import os
import re
import stat
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SOURCE_HEADERS = [
    'Detalle', 'Periodo', 'Impuesto', 'Concepto', 'Subpcto', 'Descripción',
    'Fecha Movimiento', 'Debe', 'Haber', 'Saldo',
]
WORK_HEADERS = [
    'source_row_id', 'source_row_number', 'periodo_original', 'impuesto_original',
    'concepto_original', 'subpcto_original', 'descripcion_original', 'fecha_original',
    'debe_original', 'haber_original', 'saldo_original', 'debe_numerico',
    'haber_numerico', 'saldo_numerico', 'mapped_status', 'mapping_reference',
    'derived_row_id',
]
MOVE_HEADERS = [
    'derived_row_id', 'source_row_id', 'Periodo', 'Impuesto', 'Concepto',
    'Descripción', 'Fecha movimiento', 'Obligación (Debe)', 'Pago (Haber)',
    'Diferencia', 'Estado de mapeo',
]

NAVY = '1F4E78'
BLUE = '0000FF'
LIGHT_BLUE = 'D9EAF7'
LIGHT_YELLOW = 'FFF2CC'
WHITE = 'FFFFFF'
BLACK = '000000'
THIN_GRAY = Side(style='thin', color='B7B7B7')

OOXML_MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
OOXML_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
OOXML_PACKAGE_REL = 'http://schemas.openxmlformats.org/package/2006/relationships'


def apply_formula_caches(path: Path, cache_values: dict[str, dict[str, object]]) -> None:
    """Inject verified cached values without removing the formula expressions.

    openpyxl writes formula expressions but does not calculate them. This helper
    stores values already computed from the preserved source so viewers that do
    not recalculate immediately can still display the workbook correctly.
    """
    path = Path(path)
    stage = path.with_suffix(path.suffix + '.formula-cache-stage')
    if stage.exists():
        raise RuntimeError('formula_cache_stage_exists')
    with zipfile.ZipFile(path, 'r') as source_zip:
        files = {item.filename: source_zip.read(item.filename) for item in source_zip.infolist()}
        infos = {item.filename: item for item in source_zip.infolist()}

    workbook = ET.fromstring(files['xl/workbook.xml'])
    relationships = ET.fromstring(files['xl/_rels/workbook.xml.rels'])
    relationship_targets = {
        relationship.attrib['Id']: relationship.attrib['Target']
        for relationship in relationships.findall(f'{{{OOXML_PACKAGE_REL}}}Relationship')
    }
    sheet_paths = {}
    for sheet in workbook.findall(f'.//{{{OOXML_MAIN}}}sheet'):
        relationship_id = sheet.attrib[f'{{{OOXML_REL}}}id']
        target = relationship_targets[relationship_id].lstrip('/')
        full_path = target if target.startswith('xl/') else f'xl/{target}'
        sheet_paths[sheet.attrib['name']] = full_path

    for sheet_name, cells in cache_values.items():
        sheet_path = sheet_paths.get(sheet_name)
        if not sheet_path or sheet_path not in files:
            raise RuntimeError(f'formula_cache_sheet_missing:{sheet_name}')
        root = ET.fromstring(files[sheet_path])
        indexed_cells = {
            cell.attrib.get('r'): cell
            for cell in root.findall(f'.//{{{OOXML_MAIN}}}c')
        }
        for coordinate, value in cells.items():
            cell = indexed_cells.get(coordinate)
            if cell is None or cell.find(f'{{{OOXML_MAIN}}}f') is None:
                raise RuntimeError(f'formula_cache_cell_missing:{sheet_name}!{coordinate}')
            cached = cell.find(f'{{{OOXML_MAIN}}}v')
            if cached is None:
                cached = ET.SubElement(cell, f'{{{OOXML_MAIN}}}v')
            if isinstance(value, str):
                cell.set('t', 'str')
                cached.text = value
            else:
                if cell.attrib.get('t') == 'str':
                    del cell.attrib['t']
                cached.text = str(value)
        files[sheet_path] = ET.tostring(root, encoding='utf-8', xml_declaration=True)

    try:
        with zipfile.ZipFile(stage, 'w') as destination_zip:
            for name, data in files.items():
                destination_zip.writestr(infos[name], data)
        os.chmod(stage, 0o600)
        os.replace(stage, path)
        os.chmod(path, 0o600)
    except Exception:
        if stage.exists():
            os.chmod(stage, 0o600)
        raise


def parse_amount(value: str) -> float | None:
    raw = str(value or '').strip()
    if not raw or raw in {'-', '—'}:
        return 0.0
    negative = raw.startswith('(') and raw.endswith(')')
    raw = raw.strip('()').replace('$', '').replace(' ', '')
    # Accept explicit decimal conventions; never erase the decimal separator.
    if re.fullmatch(r'-?\d{1,3}(?:,\d{3})+\.\d{2}', raw):
        raw = raw.replace(',', '')
    elif re.fullmatch(r'-?\d{1,3}(?:\.\d{3})+,\d{2}', raw):
        raw = raw.replace('.', '').replace(',', '.')
    elif re.fullmatch(r'-?\d+(?:[.,]\d{2})?', raw):
        raw = raw.replace(',', '.')
    else:
        return None
    if not re.fullmatch(r'-?\d+(?:\.\d+)?', raw):
        return None
    amount = float(raw)
    return -abs(amount) if negative else amount


def append_literal_row(sheet, values):
    """Only code-generated formulas may execute; source text stays literal."""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def set_table_style(ws, header_row: int, widths: dict[int, float], freeze: str = 'A2'):
    header_fill = PatternFill('solid', fgColor=NAVY)
    header_font = Font(name='Arial', size=10, bold=True, color=WHITE)
    body_font = Font(name='Arial', size=10, color=BLACK)
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}'
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[header_row].height = 30


def main() -> None:
    source = Path(os.environ['CCMA_SOURCE_FILE']).resolve()
    output = Path(os.environ['CCMA_WORKBOOK_FILE']).resolve()
    period_label = os.environ['CCMA_PERIOD_LABEL']
    source_sha = os.environ['CCMA_SOURCE_SHA256']
    generated_at = os.environ['CCMA_GENERATED_AT']

    source_stat = source.stat()
    if not stat.S_ISREG(source_stat.st_mode) or stat.S_IMODE(source_stat.st_mode) != 0o600:
        raise RuntimeError('source_file_not_regular_or_not_mode_600')
    with source.open('r', encoding='utf-8', newline='') as handle:
        data = list(csv.reader(handle))
    if not data or data[0] != SOURCE_HEADERS:
        raise RuntimeError('source_header_unexpected')
    rows = data[1:]
    if not rows:
        raise RuntimeError('source_has_no_movements')
    if any(len(row) > len(SOURCE_HEADERS) + 1 for row in rows):
        raise RuntimeError('source_row_width_unexpected')

    records = []
    for source_index, raw_row in enumerate(rows, start=1):
        if len(raw_row) == len(SOURCE_HEADERS) + 1:
            normalized = raw_row[-len(SOURCE_HEADERS):]
            amounts = tuple(parse_amount(normalized[column]) for column in (7, 8, 9))
            invalid_numeric = any(amount is None for amount in amounts)
            if invalid_numeric:
                # Do not publish zero-valued obligations/caches for an unknown amount.
                raise ValueError('ccma_amount_unreadable')
            status = 'unmapped'
            reference = 'Fila de 11 celdas alineada por las últimas 10 con el encabezado exportado; sin mapeo fiscal confirmado.'
            derived_id = f'DER-{source_index:04d}'
        else:
            normalized = [''] * len(SOURCE_HEADERS)
            amounts = (None, None, None)
            status = 'ambiguous'
            reference = 'Fila fuente estructural o subtotal sin las 11 celdas de un movimiento; preservada sin interpretación.'
            derived_id = ''
        records.append({
            'source_index': source_index,
            'raw': raw_row,
            'normalized': normalized,
            'amounts': amounts,
            'status': status,
            'reference': reference,
            'derived_id': derived_id,
        })

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = 'auto'
    workbook.properties.title = f'CCMA Obligaciones y Pagos {period_label}'
    workbook.properties.creator = 'Hermes Agent'
    workbook.properties.description = 'Libro read-only basado en copia local de tabla visible de CCMA.'

    # Preserve the CSV as evidence; render monetary cells numerically in Excel.
    source_ws = workbook.create_sheet('Fuente CCMA')
    source_ws.append(['Control original', *SOURCE_HEADERS])
    for row in rows:
        displayed = list(row)
        if len(row) == 11:
            for col in (8, 9, 10):
                value = parse_amount(row[col])
                if value is not None:
                    displayed[col] = value
        append_literal_row(source_ws, displayed)
    for row in source_ws.iter_rows(min_row=2, min_col=9, max_col=11):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00;[Red](#,##0.00);-'
    set_table_style(
        source_ws, 1,
        {1: 20, 2: 12, 3: 14, 4: 13, 5: 12, 6: 42, 7: 16, 8: 15, 9: 15, 10: 15},
    )
    for row in source_ws.iter_rows(min_row=2, min_col=8, max_col=10):
        for cell in row:
            cell.font = Font(name='Arial', size=10, color=BLUE)

    work_ws = workbook.create_sheet('Trabajo CCMA')
    work_ws.append(WORK_HEADERS)
    work_comment = Comment(
        'Los importes numéricos se obtuvieron mediante normalización técnica de las columnas originales. '
        'Las etiquetas y clasificación fiscal no se infieren; quedan como excepciones hasta confirmar un mapeo vigente.',
        'Hermes Agent',
    )
    for index, record in enumerate(records, start=1):
        row = record['normalized']
        amounts = record['amounts']
        source_row = record['source_index'] + 1
        source_id = f'SRC-{record["source_index"]:04d}'
        derived_id = record['derived_id']
        append_literal_row(work_ws, [
            source_id, source_row, row[1], row[2], row[3], row[4], row[5], row[6],
            *(amounts[i] if amounts[i] is not None else row[7+i] for i in range(3)),
            amounts[0], amounts[1], amounts[2],
            record['status'], record['reference'], derived_id,
        ])
        work_ws.cell(index + 1, 12).comment = work_comment
    set_table_style(
        work_ws, 1,
        {1: 15, 2: 17, 3: 14, 4: 16, 5: 14, 6: 14, 7: 42, 8: 18, 9: 16, 10: 16, 11: 16, 12: 16, 13: 16, 14: 16, 15: 14, 16: 55, 17: 15},
    )
    for row in work_ws.iter_rows(min_row=2, min_col=9, max_col=14):
        for cell in row:
            cell.number_format = '#,##0.00;[Red](#,##0.00);-'
            cell.font = Font(name='Arial', size=10, color=BLUE)
    for row in work_ws.iter_rows(min_row=2, min_col=15, max_col=16):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor=LIGHT_YELLOW)

    movement_ws = workbook.create_sheet('Obligaciones y Pagos')
    movement_ws.append(MOVE_HEADERS)
    movement_records = [
        (work_index, record)
        for work_index, record in enumerate(records, start=1)
        if record['derived_id']
    ]
    for index, (work_index, record) in enumerate(movement_records, start=1):
        work_row = work_index + 1
        movement_ws.append([
            f"='Trabajo CCMA'!Q{work_row}",
            f"='Trabajo CCMA'!A{work_row}",
            f"='Trabajo CCMA'!C{work_row}",
            f"='Trabajo CCMA'!D{work_row}",
            f"='Trabajo CCMA'!E{work_row}",
            f"='Trabajo CCMA'!G{work_row}",
            f"='Trabajo CCMA'!H{work_row}",
            f"=IFERROR('Trabajo CCMA'!L{work_row},0)",
            f"=IFERROR('Trabajo CCMA'!M{work_row},0)",
            f'=H{index + 1}-I{index + 1}',
            f"='Trabajo CCMA'!O{work_row}",
        ])
    set_table_style(
        movement_ws, 1,
        {1: 15, 2: 15, 3: 14, 4: 16, 5: 14, 6: 42, 7: 18, 8: 18, 9: 18, 10: 18, 11: 15},
    )
    for row in movement_ws.iter_rows(min_row=2, min_col=8, max_col=10):
        for cell in row:
            cell.number_format = '#,##0.00;[Red](#,##0.00);-'
            cell.font = Font(name='Arial', size=10, color=BLACK)

    summary_ws = workbook.create_sheet('Resumen')
    summary_ws['A1'] = f'CCMA — Obligaciones y Pagos — {period_label}'
    summary_ws['A1'].font = Font(name='Arial', size=14, bold=True, color=WHITE)
    summary_ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    summary_ws.merge_cells('A1:B1')
    summary_ws['A3'] = 'Indicador'
    summary_ws['B3'] = 'Valor'
    movement_end_row = max(2, len(movement_records) + 1)
    summary_rows = [
        ('Movimientos fuente', f"=COUNTA('Trabajo CCMA'!A2:A{len(records) + 1})"),
        ('Obligaciones (Debe)', f"=SUM('Obligaciones y Pagos'!H2:H{movement_end_row})"),
        ('Pagos (Haber)', f"=SUM('Obligaciones y Pagos'!I2:I{movement_end_row})"),
        ('Diferencia (Debe - Haber)', '=B5-B6'),
        ('Excepciones sin mapeo confirmado', f"=COUNTIF('Trabajo CCMA'!O2:O{len(records) + 1},\"unmapped\")+COUNTIF('Trabajo CCMA'!O2:O{len(records) + 1},\"ambiguous\")"),
        ('Estado de fórmulas', 'Pendiente de recalcular con LibreOffice Calc o Microsoft Excel'),
    ]
    for row_number, row in enumerate(summary_rows, start=4):
        summary_ws.cell(row_number, 1, row[0])
        summary_ws.cell(row_number, 2, row[1])
    set_table_style(summary_ws, 3, {1: 48, 2: 32}, 'A4')
    for cell in ('B5', 'B6', 'B7'):
        summary_ws[cell].number_format = '#,##0.00;[Red](#,##0.00);-'
    summary_ws['B9'].fill = PatternFill('solid', fgColor=LIGHT_YELLOW)
    summary_ws['B9'].alignment = Alignment(wrap_text=True, vertical='center')
    summary_ws.row_dimensions[9].height = 42

    exceptions_ws = workbook.create_sheet('Excepciones')
    exceptions_ws.append(['source_row_id', 'derived_row_id', 'Período', 'Descripción', 'Estado', 'Motivo'])
    for index in range(1, len(records) + 1):
        work_row = index + 1
        exceptions_ws.append([
            f"='Trabajo CCMA'!A{work_row}",
            f"='Trabajo CCMA'!Q{work_row}",
            f"='Trabajo CCMA'!C{work_row}",
            f"='Trabajo CCMA'!G{work_row}",
            f"='Trabajo CCMA'!O{work_row}",
            f"='Trabajo CCMA'!P{work_row}",
        ])
    set_table_style(exceptions_ws, 1, {1: 15, 2: 15, 3: 14, 4: 45, 5: 16, 6: 62})
    for row in exceptions_ws.iter_rows(min_row=2):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor=LIGHT_YELLOW)

    meta_ws = workbook.create_sheet('Trazabilidad')
    metadata = [
        ('Campo', 'Valor'),
        ('Período solicitado', period_label),
        ('Fuente copiada', source.name),
        ('Hash SHA-256 de fuente', source_sha),
        ('Filas de fuente', len(rows)),
        ('Fecha/hora de copia procesada', generated_at),
        ('Origen', 'Consulta CCMA read-only a través de ARCA; tabla visible copiada a CSV local'),
        ('Integridad de fuente', 'Hash registrado antes de construir el libro; verificar nuevamente al cierre.'),
        ('Mapeos fiscales', 'No se aplicó mapeo candidato sin confirmación vigente; todas las filas quedan como excepción.'),
        ('Fórmulas', 'Incluidas con recálculo automático al abrir; LibreOffice Calc no estaba disponible durante esta generación.'),
        ('Acciones excluidas', 'No se generó VEP, pago, reimputación, compensación ni presentación.'),
    ]
    for row in metadata:
        append_literal_row(meta_ws, row)
    set_table_style(meta_ws, 1, {1: 36, 2: 98})
    for row in meta_ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical='top')
    meta_ws['B4'].font = Font(name='Arial', size=10, color=BLUE)

    formula_caches = {
        'Obligaciones y Pagos': {},
        'Resumen': {},
        'Excepciones': {},
    }
    total_debe = 0.0
    total_haber = 0.0
    for movement_index, (work_index, record) in enumerate(movement_records, start=1):
        row_number = movement_index + 1
        normalized = record['normalized']
        debe = record['amounts'][0] if record['amounts'][0] is not None else 0.0
        haber = record['amounts'][1] if record['amounts'][1] is not None else 0.0
        total_debe += debe
        total_haber += haber
        formula_caches['Obligaciones y Pagos'].update({
            f'A{row_number}': record['derived_id'],
            f'B{row_number}': f"SRC-{record['source_index']:04d}",
            f'C{row_number}': normalized[1],
            f'D{row_number}': normalized[2],
            f'E{row_number}': normalized[3],
            f'F{row_number}': normalized[5],
            f'G{row_number}': normalized[6],
            f'H{row_number}': debe,
            f'I{row_number}': haber,
            f'J{row_number}': debe - haber,
            f'K{row_number}': record['status'],
        })

    exception_count = 0
    for work_index, record in enumerate(records, start=1):
        row_number = work_index + 1
        if record['status'] in {'unmapped', 'ambiguous'}:
            exception_count += 1
        formula_caches['Excepciones'].update({
            f'A{row_number}': f"SRC-{record['source_index']:04d}",
            f'B{row_number}': record['derived_id'] or '',
            f'C{row_number}': record['normalized'][1],
            f'D{row_number}': record['normalized'][5],
            f'E{row_number}': record['status'],
            f'F{row_number}': record['reference'],
        })

    formula_caches['Resumen'] = {
        'B4': len(records),
        'B5': total_debe,
        'B6': total_haber,
        'B7': total_debe - total_haber,
        'B8': exception_count,
    }

    workbook.save(output)
    apply_formula_caches(output, formula_caches)
    os.chmod(output, 0o600)
    print(f'workbook_created={output.name}')
    print(f'source_rows={len(rows)}')
    print(f'exceptions={len(rows)}')
    print('formula_recalculation=deferred_no_libreoffice')


if __name__ == '__main__':
    main()
