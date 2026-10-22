import tempfile
import unittest
import importlib.util
from datetime import date
from pathlib import Path

from plugins.platforms.telegram.vencimientos_export import generate


ROWS = [{
    "id_impuesto": 30,
    "impuesto": "IVA",
    "id_concepto": 19,
    "concepto": "Declaración jurada",
    "periodo": "2026-08",
    "anticipo_cuota": "0",
    "tipo": "PRESENTACION",
    "fecha": "2026-09-18",
    "formularios": "F. 2051",
    "estado": "pendiente",
}]


class VencimientosExportTests(unittest.TestCase):
    def test_pair_same_concept_keeps_two_operations_and_human_labels(self):
        from openpyxl import load_workbook
        rows=[dict(ROWS[0],tipo='PAGO'),ROWS[0]]
        with tempfile.TemporaryDirectory() as directory:
            excel=Path(directory)/'cliente-demo-vencimientos-arca.xlsx'
            ics=Path(directory)/'cliente-demo-vencimientos-arca.ics'
            generate('xlsx',excel,9,'cliente-demo',rows)
            generate('ics',ics,9,'cliente-demo',rows)
            book=load_workbook(excel,data_only=True)
            self.assertEqual(book.active.max_column,10)
            self.assertEqual([book.active.cell(n,5).value for n in (2,3)],['Pago','Presentación'])
            book.close()
            text=ics.read_text()
            for label in ('Pago','Presentación'):
                self.assertIn(f'SUMMARY:{label} · IVA · Declaración jurada',text)
            uids=[x for x in text.splitlines() if x.startswith('UID:')]
            self.assertEqual(len(set(uids)),2)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is provided by the fiscal runtime")
    def test_excel_is_readable_and_preserves_the_operational_columns(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "cliente-demo-vencimientos-arca.xlsx"
            generate("xlsx", output, 9, "cliente-demo", ROWS)
            book = load_workbook(output, data_only=False)
            sheet = book["Vencimientos"]
            self.assertEqual(
                [cell.value for cell in sheet[1]],
                ["Impuesto", "Concepto", "Período", "Anticipo/Cuota",
                 "Tipo de operación", "Vencimiento", "Formularios", "Estado", "Organismo", "Fuente vigente"],
            )
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(str(sheet["F2"].value)[:10], date(2026, 9, 18).isoformat())
            self.assertEqual(sheet["F2"].number_format, "dd/mm/yyyy")
            self.assertEqual(sheet['I2'].value,'ARCA')
            self.assertEqual(sheet['J2'].value,'Sí')
            self.assertEqual(output.stat().st_mode & 0o777, 0o600)
            book.close()

    def test_ics_is_importable_as_an_all_day_calendar_and_has_stable_uid(self):
        with tempfile.TemporaryDirectory() as directory:
            first = Path(directory) / "cliente-demo-vencimientos-arca.ics"
            generate("ics", first, 9, "cliente-demo", ROWS)
            content = first.read_text(encoding="utf-8")
            self.assertIn("BEGIN:VCALENDAR", content)
            self.assertIn("DTSTART;VALUE=DATE:20260918", content)
            self.assertIn("DTEND;VALUE=DATE:20260919", content)
            self.assertIn("SUMMARY:Presentación · IVA · Declaración jurada", content)
            uid = next(line for line in content.splitlines() if line.startswith("UID:"))

            second_dir = Path(directory) / "second"
            second_dir.mkdir()
            second = second_dir / "cliente-demo-vencimientos-arca.ics"
            generate("ics", second, 9, "cliente-demo", ROWS)
            self.assertIn(uid, second.read_text(encoding="utf-8"))
            self.assertEqual(first.stat().st_mode & 0o777, 0o600)

    def test_rejects_unsafe_slug_and_existing_destination(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "../cliente-vencimientos-arca.ics"
            with self.assertRaisesRegex(ValueError, "invalid_export_request"):
                generate("ics", output, 9, "../cliente", ROWS)
            existing = Path(directory) / "cliente-demo-vencimientos-arca.ics"
            existing.write_text("keep", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "invalid_export_target"):
                generate("ics", existing, 9, "cliente-demo", ROWS)
            self.assertEqual(existing.read_text(encoding="utf-8"), "keep")


if __name__ == "__main__":
    unittest.main()
